"""
Asistente IA flotante: preguntas libres con datos del sistema y Excel.

La IA se simula: se prueba lo que el sistema le envia y lo que hace con la
respuesta, no al modelo.
"""
import json
from io import BytesIO
from unittest.mock import patch

from django.core.cache import cache
from django.test import override_settings
from openpyxl import load_workbook

from apps.ia.asistente import (
    INSTRUCCION_ASISTENTE, MAX_PREGUNTAS_POR_MINUTO, construir_contexto_sistema, limpiar_respuesta,
)
from apps.ia.models import ConsultaIA
from apps.ia.services import ErrorIA

from .base import BaseSigmap

URL = '/ai/asistente/preguntar/'

SALIDA_IA = {
    'respuesta': 'El **pH** de Estanque de Prueba está alto.',
    'hallazgos': [{'piscina': 'Estanque de Prueba', 'parametro': 'pH',
                   'descripcion': 'pH 9.5 fuera del rango 6.5-8.5', 'gravedad': 'alta'}],
    'tablas': [{'titulo': 'Últimas lecturas', 'columnas': ['Parámetro', 'Valor'],
                'filas': [['pH', '9.5'], ['Oxígeno disuelto', '6.2']]}],
    'recomendaciones': ['Hacer recambio parcial de agua.'],
    'enlaces': [{'etiqueta': 'Ver alertas', 'url': '/alertas/'}],
}


def _post(cliente, pregunta):
    return cliente.post(URL, data=json.dumps({'pregunta': pregunta}), content_type='application/json')


class BaseAsistente(BaseSigmap):
    def setUp(self):
        cache.clear()
        for valor in (7.2, 9.5):
            self.enviar_lectura('SEN-PH', valor)
        self.enviar_lectura('SEN-OX', 6.2)
        self.client.force_login(self.instructor)


class ContextoSistemaTest(BaseAsistente):
    """Lo que la IA recibe del software."""

    def test_incluye_piscinas_activas_con_sus_lecturas(self):
        ctx = construir_contexto_sistema()
        piscina = ctx['piscinas_activas'][0]
        self.assertEqual(piscina['codigo'], 'GEO-TEST')
        valores = {s['sensor']: s['ultimo_valor'] for s in piscina['sensores']}
        self.assertEqual(valores['SEN-PH'], 9.5)
        self.assertTrue(piscina['estadisticas_7d'])

    def test_incluye_alertas_y_modulos(self):
        ctx = construir_contexto_sistema()
        self.assertTrue(ctx['alertas_activas'])           # el 9.5 dispara alerta
        self.assertIn('/historial/', [m['url'] for m in ctx['modulos']])

    def test_es_serializable_y_sin_correos(self):
        texto = json.dumps(construir_contexto_sistema(), ensure_ascii=False)
        self.assertNotIn('@sena.edu.co', texto)


@patch('apps.ia.asistente._llamar_ia', return_value=(SALIDA_IA, 'deepseek-flash'))
class PreguntarTest(BaseAsistente):
    """Camino feliz, validaciones y errores del endpoint."""

    def test_responde_y_guarda_la_consulta(self, llamada):
        respuesta = _post(self.client, '¿Cómo está el pH?')
        self.assertEqual(respuesta.status_code, 200)
        datos = respuesta.json()['data']
        self.assertEqual(datos['hallazgos'][0]['gravedad'], 'alta')
        self.assertEqual(datos['modelo'], 'deepseek-flash')
        self.assertEqual(datos['url_excel'], f'/ai/asistente/{datos["id"]}/excel/')
        self.assertTrue(ConsultaIA.objects.filter(pk=datos['id'], usuario=self.instructor).exists())

    def test_envia_datos_pregunta_e_instruccion(self, llamada):
        _post(self.client, '¿Cómo está el pH?')
        prompt = llamada.call_args.args[0]
        self.assertIn('DATOS_DEL_SISTEMA', prompt)
        self.assertIn('GEO-TEST', prompt)
        self.assertIn('¿Cómo está el pH?', prompt)
        self.assertEqual(llamada.call_args.kwargs['instruccion'], INSTRUCCION_ASISTENTE)

    def test_recuerda_la_pregunta_anterior(self, llamada):
        _post(self.client, '¿Cómo está el pH?')
        _post(self.client, '¿Y en la otra piscina?')
        self.assertIn('¿Cómo está el pH?', llamada.call_args.args[0].split('PREGUNTA_DEL_USUARIO')[0])

    def test_pregunta_vacia_es_400(self, llamada):
        self.assertEqual(_post(self.client, '   ').status_code, 400)
        llamada.assert_not_called()

    def test_pregunta_demasiado_larga_es_400(self, llamada):
        self.assertEqual(_post(self.client, 'x' * 501).status_code, 400)

    def test_cuerpo_invalido_es_400(self, llamada):
        respuesta = self.client.post(URL, data='no-json', content_type='application/json')
        self.assertEqual(respuesta.status_code, 400)

    def test_limite_de_preguntas_por_minuto(self, llamada):
        for _ in range(MAX_PREGUNTAS_POR_MINUTO):
            self.assertEqual(_post(self.client, 'hola').status_code, 200)
        self.assertEqual(_post(self.client, 'hola').status_code, 429)

    def test_sin_sesion_es_401(self, llamada):
        self.client.logout()
        self.assertEqual(_post(self.client, 'hola').status_code, 401)

    def test_solo_post(self, llamada):
        self.assertEqual(self.client.get(URL).status_code, 405)


class IaCaidaTest(BaseAsistente):
    @patch('apps.ia.asistente._llamar_ia', side_effect=ErrorIA('Sin conexión con DeepSeek'))
    def test_fallo_de_la_ia_es_503_con_mensaje(self, llamada):
        respuesta = _post(self.client, '¿Cómo está el pH?')
        self.assertEqual(respuesta.status_code, 503)
        self.assertFalse(respuesta.json()['success'])
        self.assertFalse(ConsultaIA.objects.exists())


class LimpiezaTest(BaseAsistente):
    """Nunca se confía en la forma de la salida del modelo."""

    def test_descarta_enlaces_externos(self):
        limpia = limpiar_respuesta({'respuesta': 'x', 'enlaces': [
            {'etiqueta': 'malo', 'url': 'https://evil.example'},
            {'etiqueta': 'bueno', 'url': '/alertas/'}]})
        self.assertEqual([e['url'] for e in limpia['enlaces']], ['/alertas/'])

    def test_normaliza_gravedad_y_tipos(self):
        limpia = limpiar_respuesta({'respuesta': 5, 'hallazgos': [
            {'descripcion': 'algo', 'gravedad': 'extrema'}, 'basura'],
            'tablas': [{'titulo': 'T', 'columnas': ['A', 'B'], 'filas': [['1'], 'mala']}]})
        self.assertEqual(limpia['hallazgos'][0]['gravedad'], 'media')
        self.assertEqual(limpia['tablas'][0]['filas'], [['1', '']])

    def test_salida_vacia_no_rompe(self):
        self.assertTrue(limpiar_respuesta(None)['respuesta'])


@patch('apps.ia.asistente._llamar_ia', return_value=(SALIDA_IA, 'deepseek-flash'))
class ExcelTest(BaseAsistente):
    """Botón de descarga del informe."""

    def _consulta(self, cliente):
        return _post(cliente, '¿Cómo está el pH?').json()['data']['id']

    def test_excel_con_hojas_de_respuesta_tablas_y_evidencia(self, llamada):
        pk = self._consulta(self.client)
        respuesta = self.client.get(f'/ai/asistente/{pk}/excel/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('.xlsx', respuesta['Content-Disposition'])
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ['Respuesta', 'Hallazgos', 'Últimas lecturas', 'Datos del sistema'])
        self.assertEqual(libro['Últimas lecturas']['B2'].value, 9.5)   # número real, no texto

    def test_otro_usuario_no_descarga_la_consulta_ajena(self, llamada):
        pk = self._consulta(self.client)
        self.client.force_login(self.operario)
        self.assertEqual(self.client.get(f'/ai/asistente/{pk}/excel/').status_code, 404)

    def test_el_instructor_descarga_la_de_otro(self, llamada):
        self.client.force_login(self.operario)
        pk = self._consulta(self.client)
        self.client.force_login(self.instructor)
        self.assertEqual(self.client.get(f'/ai/asistente/{pk}/excel/').status_code, 200)


class ContextoGlobalTest(BaseAsistente):
    """El círculo flotante recibe la URL y si la IA está disponible."""

    @override_settings(DEEPSEEK_API_KEY='clave', GEMINI_API_KEY='')
    def test_expone_url_e_ia_disponible(self):
        respuesta = self.client.get('/historial/')
        datos = respuesta.context['agente_ia']
        self.assertEqual(datos['url_preguntar'], URL)
        self.assertTrue(datos['ia_disponible'])


class RobustezTest(BaseAsistente):
    """El chat nunca recibe una página 500: siempre JSON con un mensaje útil."""

    @patch('apps.ia.asistente._llamar_ia', return_value=(SALIDA_IA, 'deepseek-flash'))
    def test_sin_tabla_de_consultas_responde_igual_sin_excel(self, llamada):
        from django.db import DatabaseError
        with patch('apps.ia.models.ConsultaIA.save', side_effect=DatabaseError('no existe consultas_ia')):
            respuesta = _post(self.client, '¿Qué piscinas hay activas?')
        self.assertEqual(respuesta.status_code, 200)
        datos = respuesta.json()['data']
        self.assertIsNone(datos['url_excel'])
        self.assertIn('migrate', datos['aviso'])

    @patch('apps.ia.asistente.construir_contexto_sistema', side_effect=RuntimeError('boom'))
    def test_error_inesperado_es_json_500(self, contexto):
        respuesta = _post(self.client, '¿Qué piscinas hay activas?')
        self.assertEqual(respuesta.status_code, 500)
        self.assertFalse(respuesta.json()['success'])
        self.assertIn('logs/ia.log', respuesta.json()['message'])

    def test_una_seccion_rota_no_tumba_el_contexto(self):
        with patch('apps.ia.asistente._reportes', side_effect=RuntimeError('tabla rota')):
            ctx = construir_contexto_sistema()
        self.assertEqual(ctx['reportes_recientes'], [])
        self.assertTrue(ctx['piscinas_activas'])

    def test_lista_todas_las_piscinas_registradas(self):
        ctx = construir_contexto_sistema()
        self.assertIn('GEO-TEST', [p['codigo'] for p in ctx['piscinas_registradas']])
