"""
Exportacion del historial a Excel y PDF.
"""
from io import BytesIO

from openpyxl import load_workbook

from apps.monitoreo.models import Lectura

from .base import BaseSigmap


class ExportacionTest(BaseSigmap):
    """Las dos salidas respetan filtros y permisos."""

    def setUp(self):
        for valor in (7.1, 7.4, 9.5):          # 9.5 queda critico
            self.enviar_lectura('SEN-PH', valor)
        self.enviar_lectura('SEN-OX', 6.2)
        self.client.force_login(self.instructor)

    def _excel(self, consulta=''):
        respuesta = self.client.get('/historial/excel/' + consulta)
        self.assertEqual(respuesta.status_code, 200)
        return respuesta, load_workbook(BytesIO(respuesta.content))

    def test_excel_tiene_resumen_y_lecturas(self):
        respuesta, libro = self._excel()
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        self.assertIn('.xlsx', respuesta['Content-Disposition'])
        self.assertEqual(libro.sheetnames, ['Resumen', 'Lecturas'])
        self.assertEqual(libro['Lecturas'].max_row, 1 + lecturas_totales())

    def test_excel_guarda_valores_y_fechas_como_datos(self):
        """Valores numericos reales: el usuario puede sumar, filtrar y graficar."""
        _, libro = self._excel()
        hoja = libro['Lecturas']
        self.assertIsInstance(hoja['E2'].value, float)
        self.assertEqual(hoja['A2'].number_format, 'dd/mm/yyyy hh:mm:ss')
        self.assertEqual(hoja.freeze_panes, 'A2')
        self.assertTrue(hoja.auto_filter.ref)

    def test_excel_respeta_el_filtro_de_parametro(self):
        _, libro = self._excel(f'?parametro={self.ph.pk}')
        self.assertEqual(libro['Lecturas'].max_row, 1 + 3)

    def test_resumen_por_parametro(self):
        _, libro = self._excel()
        textos = [c.value for fila in libro['Resumen'].iter_rows() for c in fila]
        self.assertIn('pH', textos)
        self.assertIn('Oxígeno disuelto', textos)

    def test_pdf_se_genera(self):
        respuesta = self.client.get('/historial/pdf/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertTrue(respuesta.content.startswith(b'%PDF'))

    def test_pdf_sin_datos_no_falla(self):
        """Caso limite: un rango sin lecturas genera el informe vacio, no un error."""
        respuesta = self.client.get('/historial/pdf/?desde=2020-01-01&hasta=2020-01-02')
        self.assertEqual(respuesta.status_code, 200)

    def test_requiere_sesion(self):
        self.client.logout()
        for ruta in ('/historial/excel/', '/historial/pdf/'):
            self.assertEqual(self.client.get(ruta).status_code, 302)

    def test_la_pantalla_muestra_los_botones(self):
        respuesta = self.client.get('/historial/')
        self.assertContains(respuesta, 'Exportar Excel')
        self.assertContains(respuesta, 'Exportar PDF')


def lecturas_totales():
    return Lectura.objects.count()
