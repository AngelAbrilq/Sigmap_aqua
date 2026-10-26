"""
Exportacion del historial de lecturas a Excel (.xlsx) y PDF.

Las dos salidas parten del mismo contexto (filtros + resumen + filas), asi
el Excel y el PDF siempre muestran los mismos numeros que la pantalla.

- Excel: hoja "Resumen" por parametro y hoja "Lecturas" con fechas y valores
  como datos reales (se pueden filtrar, ordenar y graficar), colores por
  estado, filtros automaticos y encabezado fijo.
- PDF: informe horizontal A4 con logo, filtros aplicados, resumen por
  parametro y la tabla de lecturas paginada y numerada.
"""
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.db.models import Avg, Count, Max, Min, Q
from django.utils import timezone

from apps.piscinas.models import Geomembrana

from .models import Lectura, TipoParametro

# Un PDF de miles de paginas no lo lee nadie: se acota y se avisa.
MAX_FILAS_PDF = 1500

COLOR_PRIMARIO = '1E3A8A'
COLORES_ESTADO = {            # (fondo, texto) por estado de lectura
    'normal': ('DCFCE7', '166534'),
    'riesgo': ('FEF3C7', '92400E'),
    'critico': ('FEE2E2', '991B1B'),
}
RUTA_LOGO = Path(settings.BASE_DIR) / 'static' / 'img' / 'Logo.jpeg'


@dataclass
class ContextoExportacion:
    """Todo lo que necesitan el Excel y el PDF."""
    filtros: list            # [(etiqueta, valor legible)]
    resumen_total: dict
    resumen_parametros: list  # dicts por parametro
    lecturas: list           # instancias de Lectura (acotadas)
    total_filas: int
    usuario: str
    generado: datetime = field(default_factory=timezone.localtime)


def construir_contexto(filtros, queryset, usuario, limite):
    """
    Reune filtros legibles, resumenes y filas para exportar.

    :param filtros: dict devuelto por _filtrar_lecturas
    :param queryset: QuerySet de Lectura ya filtrado
    :param usuario: usuarios.Usuario que exporta
    :param limite: maximo de filas de detalle
    :return: ContextoExportacion
    """
    piscina = Geomembrana.objects.filter(pk=filtros['filtro_piscina']).first() \
        if str(filtros['filtro_piscina']).isdigit() else None
    parametro = TipoParametro.objects.filter(pk=filtros['filtro_parametro']).first() \
        if str(filtros['filtro_parametro']).isdigit() else None
    estado = dict(Lectura.ESTADOS_LECTURA).get(filtros['filtro_estado'], 'Todos')

    filtros_legibles = [
        ('Desde', filtros['desde']),
        ('Hasta', filtros['hasta']),
        ('Piscina', piscina.nombre_piscina if piscina else 'Todas'),
        ('Parámetro', parametro.nombre_parametro if parametro else 'Todos'),
        ('Estado', estado),
    ]

    agregados = dict(
        total=Count('id'),
        fuera_rango=Count('id', filter=Q(dentro_rango=False)),
        criticas=Count('id', filter=Q(estado_lectura='critico')),
        riesgo=Count('id', filter=Q(estado_lectura='riesgo')),
        promedio=Avg('valor_medida'),
        minimo=Min('valor_medida'),
        maximo=Max('valor_medida'),
    )
    resumen_total = queryset.aggregate(**{k: v for k, v in agregados.items()
                                          if k in ('total', 'fuera_rango', 'criticas', 'riesgo')})
    resumen_parametros = list(
        queryset.order_by()
        .values('tipo_parametro__nombre_parametro', 'tipo_parametro__unidad_medida')
        .annotate(**agregados)
        .order_by('tipo_parametro__nombre_parametro')
    )
    for fila in resumen_parametros:
        fila['pct_fuera'] = (fila['fuera_rango'] / fila['total'] * 100) if fila['total'] else 0

    lecturas = list(queryset.order_by('-timestamp_lectura')[:limite])
    return ContextoExportacion(
        filtros=filtros_legibles,
        resumen_total=resumen_total,
        resumen_parametros=resumen_parametros,
        lecturas=lecturas,
        total_filas=resumen_total['total'] or 0,
        usuario=getattr(usuario, 'nombre_completo', '') or getattr(usuario, 'email', ''),
    )


def _num(valor, decimales=2):
    """Decimal/None a float redondeado, para que Excel lo trate como numero."""
    return round(float(valor), decimales) if valor is not None else None


# ----------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------
def generar_excel(ctx):
    """
    Libro .xlsx con hojas Resumen y Lecturas.

    :param ctx: ContextoExportacion
    :return: bytes del archivo
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    titulo = Font(bold=True, size=14, color=COLOR_PRIMARIO)
    negrita = Font(bold=True)
    tenue = Font(color='6B7280', size=9)
    encabezado_font = Font(bold=True, color='FFFFFF')
    encabezado_fill = PatternFill('solid', fgColor=COLOR_PRIMARIO)
    borde = Border(bottom=Side(style='thin', color='E5E7EB'))
    centro = Alignment(horizontal='center', vertical='center')

    libro = Workbook()

    # --- Hoja Resumen ---
    hoja = libro.active
    hoja.title = 'Resumen'
    hoja['A1'] = 'SIGMAP-AQUA · Historial de datos'
    hoja['A1'].font = titulo
    hoja['A2'] = (f'Generado el {ctx.generado:%d/%m/%Y %H:%M} por {ctx.usuario}')
    hoja['A2'].font = tenue

    fila = 4
    for etiqueta, valor in ctx.filtros:
        hoja.cell(fila, 1, etiqueta).font = negrita
        hoja.cell(fila, 2, valor)
        fila += 1

    fila += 1
    totales = ctx.resumen_total
    for etiqueta, valor in (('Lecturas', totales['total']), ('En riesgo', totales['riesgo']),
                            ('Críticas', totales['criticas']), ('Fuera de rango', totales['fuera_rango'])):
        hoja.cell(fila, 1, etiqueta).font = negrita
        hoja.cell(fila, 2, valor or 0)
        fila += 1

    fila += 1
    columnas = ['Parámetro', 'Unidad', 'Lecturas', 'Promedio', 'Mínimo', 'Máximo',
                'En riesgo', 'Críticas', '% fuera de rango']
    for col, texto in enumerate(columnas, 1):
        celda = hoja.cell(fila, col, texto)
        celda.font, celda.fill, celda.alignment = encabezado_font, encabezado_fill, centro
    for p in ctx.resumen_parametros:
        fila += 1
        valores = [p['tipo_parametro__nombre_parametro'], p['tipo_parametro__unidad_medida'],
                   p['total'], _num(p['promedio']), _num(p['minimo']), _num(p['maximo']),
                   p['riesgo'], p['criticas'], p['pct_fuera'] / 100]
        for col, valor in enumerate(valores, 1):
            celda = hoja.cell(fila, col, valor)
            celda.border = borde
            if col in (4, 5, 6):
                celda.number_format = '0.00'
            if col == 9:
                celda.number_format = '0.0%'
    for col, ancho in enumerate((22, 12, 11, 11, 11, 11, 11, 11, 17), 1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho

    # --- Hoja Lecturas ---
    hoja = libro.create_sheet('Lecturas')
    columnas = ['Fecha y hora', 'Piscina', 'Sensor', 'Parámetro', 'Valor', 'Unidad',
                'Estado', 'Dentro de rango']
    for col, texto in enumerate(columnas, 1):
        celda = hoja.cell(1, col, texto)
        celda.font, celda.fill, celda.alignment = encabezado_font, encabezado_fill, centro

    rellenos = {k: (PatternFill('solid', fgColor=f), Font(bold=True, color=t))
                for k, (f, t) in COLORES_ESTADO.items()}
    for i, lectura in enumerate(ctx.lecturas, 2):
        momento = timezone.localtime(lectura.timestamp_lectura).replace(tzinfo=None)
        valores = [momento, lectura.geomembrana.nombre_piscina, lectura.sensor.codigo_hardware,
                   lectura.tipo_parametro.nombre_parametro, _num(lectura.valor_medida, 4),
                   lectura.tipo_parametro.unidad_medida, lectura.get_estado_lectura_display(),
                   'Sí' if lectura.dentro_rango else 'No']
        for col, valor in enumerate(valores, 1):
            hoja.cell(i, col, valor)
        hoja.cell(i, 1).number_format = 'dd/mm/yyyy hh:mm:ss'
        hoja.cell(i, 5).number_format = '0.00'
        estilo = rellenos.get(lectura.estado_lectura)
        if estilo:
            hoja.cell(i, 7).fill, hoja.cell(i, 7).font = estilo
        hoja.cell(i, 7).alignment = centro
        hoja.cell(i, 8).alignment = centro

    for col, ancho in enumerate((20, 22, 18, 20, 10, 10, 12, 16), 1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.freeze_panes = 'A2'
    ultima = max(1, len(ctx.lecturas) + 1)
    hoja.auto_filter.ref = f'A1:H{ultima}'
    if len(ctx.lecturas) < ctx.total_filas:
        nota = hoja.cell(ultima + 2, 1, f'Se exportaron las {len(ctx.lecturas)} lecturas más '
                                        f'recientes de {ctx.total_filas}. Acota el rango de fechas '
                                        'para ver el resto.')
        nota.font = tenue

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def generar_pdf(ctx):
    """
    Informe PDF horizontal con resumen y tabla de lecturas.

    :param ctx: ContextoExportacion
    :return: bytes del archivo
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    azul = colors.HexColor('#' + COLOR_PRIMARIO)
    gris = colors.HexColor('#6B7280')
    borde = colors.HexColor('#E5E7EB')
    alterno = colors.HexColor('#F8FAFC')

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle('titulo', parent=estilos['Title'], textColor=azul,
                            fontSize=17, alignment=0, spaceAfter=2)
    subtitulo = ParagraphStyle('sub', parent=estilos['Normal'], textColor=gris, fontSize=8.5)
    seccion = ParagraphStyle('seccion', parent=estilos['Heading3'], textColor=azul,
                             fontSize=11, spaceBefore=10, spaceAfter=4)
    celda = ParagraphStyle('celda', parent=estilos['Normal'], fontSize=8, leading=10)

    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=14 * mm,
                            rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title='Historial de datos - SIGMAP-AQUA', author=ctx.usuario)
    elementos = []

    # Encabezado con logo
    texto = [Paragraph('Historial de datos', titulo),
             Paragraph(f'SIGMAP-AQUA · Generado el {ctx.generado:%d/%m/%Y %H:%M} por {ctx.usuario}',
                       subtitulo)]
    if RUTA_LOGO.exists():
        encabezado = Table([[Image(str(RUTA_LOGO), 16 * mm, 16 * mm), texto]],
                           colWidths=[20 * mm, None])
        encabezado.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                        ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
        elementos.append(encabezado)
    else:
        elementos.extend(texto)

    # Filtros y totales en una franja
    t = ctx.resumen_total
    franja = [[f'{e}: {v}' for e, v in ctx.filtros] +
              [f'Lecturas: {t["total"] or 0}', f'Críticas: {t["criticas"] or 0}']]
    franja = Table(franja)
    franja.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 8), ('TEXTCOLOR', (0, 0), (-1, -1), gris),
        ('BACKGROUND', (0, 0), (-1, -1), alterno), ('BOX', (0, 0), (-1, -1), 0.5, borde),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elementos += [Spacer(1, 6), franja]

    def tabla(datos, anchos, estado_col=None):
        """Tabla con el estilo del informe; colorea la columna de estado."""
        tb = Table(datos, colWidths=anchos, repeatRows=1)
        estilo = [
            ('BACKGROUND', (0, 0), (-1, 0), azul), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, alterno]),
            ('LINEBELOW', (0, 0), (-1, -1), 0.25, borde),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]
        if estado_col is not None:
            for fila, _ in enumerate(datos[1:], 1):
                clave = _ESTADOS_POR_ETIQUETA.get(datos[fila][estado_col])
                if clave:
                    fondo, letra = COLORES_ESTADO[clave]
                    estilo += [('BACKGROUND', (estado_col, fila), (estado_col, fila),
                                colors.HexColor('#' + fondo)),
                               ('TEXTCOLOR', (estado_col, fila), (estado_col, fila),
                                colors.HexColor('#' + letra)),
                               ('FONTNAME', (estado_col, fila), (estado_col, fila), 'Helvetica-Bold')]
        tb.setStyle(TableStyle(estilo))
        return tb

    # Resumen por parametro
    elementos.append(Paragraph('Resumen por parámetro', seccion))
    datos = [['Parámetro', 'Unidad', 'Lecturas', 'Promedio', 'Mínimo', 'Máximo',
              'En riesgo', 'Críticas', '% fuera de rango']]
    for p in ctx.resumen_parametros:
        datos.append([p['tipo_parametro__nombre_parametro'], p['tipo_parametro__unidad_medida'],
                      p['total'], f'{_num(p["promedio"]):.2f}', f'{_num(p["minimo"]):.2f}',
                      f'{_num(p["maximo"]):.2f}', p['riesgo'], p['criticas'],
                      f'{p["pct_fuera"]:.1f} %'])
    if len(datos) == 1:
        datos.append(['Sin lecturas para los filtros seleccionados'] + [''] * 8)
    elementos.append(tabla(datos, [48 * mm, 22 * mm] + [26 * mm] * 7))

    # Detalle
    elementos.append(Paragraph('Lecturas', seccion))
    if len(ctx.lecturas) < ctx.total_filas:
        elementos.append(Paragraph(
            f'Se muestran las {len(ctx.lecturas)} lecturas más recientes de {ctx.total_filas}. '
            'Para el detalle completo usa la exportación a Excel.', subtitulo))
        elementos.append(Spacer(1, 4))
    datos = [['Fecha y hora', 'Piscina', 'Sensor', 'Parámetro', 'Valor', 'Estado', 'En rango']]
    for lectura in ctx.lecturas:
        datos.append([
            timezone.localtime(lectura.timestamp_lectura).strftime('%d/%m/%Y %H:%M:%S'),
            Paragraph(lectura.geomembrana.nombre_piscina, celda),
            lectura.sensor.codigo_hardware,
            lectura.tipo_parametro.nombre_parametro,
            f'{_num(lectura.valor_medida):.2f} {lectura.tipo_parametro.unidad_medida}',
            lectura.get_estado_lectura_display(),
            'Sí' if lectura.dentro_rango else 'No',
        ])
    elementos.append(tabla(datos, [36 * mm, 52 * mm, 38 * mm, 40 * mm, 34 * mm, 26 * mm, 20 * mm],
                           estado_col=5))

    def pie(lienzo, documento):
        lienzo.saveState()
        lienzo.setFont('Helvetica', 7.5)
        lienzo.setFillColor(gris)
        lienzo.drawString(14 * mm, 8 * mm, 'SIGMAP-AQUA · Sistema Inteligente de Monitoreo Acuícola')
        lienzo.drawRightString(documento.pagesize[0] - 14 * mm, 8 * mm,
                               f'Página {lienzo.getPageNumber()}')
        lienzo.restoreState()

    doc.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()


_ESTADOS_POR_ETIQUETA = {etiqueta: clave for clave, etiqueta in Lectura.ESTADOS_LECTURA}
